"""Build .docx and .pdf versions of every sample tender, a scanned (image-only) PDF and a .pptx deck."""
import shutil
from pathlib import Path

import docx
import pymupdf

HERE = Path(__file__).resolve().parent
PUBLIC = HERE.parent / "frontend" / "public" / "samples"
FONT_CANDIDATES = [  # Devanagari-capable fonts for the Hindi sample
    "/mnt/c/Windows/Fonts/Nirmala.ttc", "C:/Windows/Fonts/Nirmala.ttc",
    "/usr/share/fonts/truetype/noto/NotoSansDevanagari-Regular.ttf",
]


def write_pdf(lines: list[str], path: Path) -> None:
    text = "\n".join(lines)
    font = next((f for f in FONT_CANDIDATES if Path(f).exists()), None)
    pdf = pymupdf.open()
    page = pdf.new_page()
    rect = pymupdf.Rect(50, 50, 545, 800)
    if font and any(ord(c) > 0x900 for c in text):
        page.insert_font(fontname="deva", fontfile=font)
        page.insert_textbox(rect, text, fontsize=10, fontname="deva")
    else:
        page.insert_textbox(rect, text, fontsize=10, fontname="helv")
    pdf.save(path)


def write_scanned(source: Path, path: Path) -> None:
    """Render a text PDF to an image and wrap it in a new PDF with no text layer."""
    with pymupdf.open(source) as src:
        pix = src[0].get_pixmap(dpi=150)
    out = pymupdf.open()
    page = out.new_page(width=pix.width * 72 / 150, height=pix.height * 72 / 150)
    page.insert_image(page.rect, pixmap=pix)
    out.save(path)


def write_pptx(source: Path, path: Path) -> None:
    """A slide deck version of a tender: a title slide, then one slide per numbered item."""
    from pptx import Presentation
    from pptx.util import Pt

    lines = [l for l in source.read_text(encoding="utf-8").splitlines() if l.strip()]
    items = [l for l in lines if l[:1].isdigit()]
    prs = Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = lines[2]
    title.placeholders[1].text = "Technical specifications (demonstration only)"
    for item in items:
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank: a slide title would leak into the item text
        box = slide.shapes.add_textbox(Pt(40), Pt(60), Pt(640), Pt(400)).text_frame
        box.word_wrap = True
        box.text = item  # keep the "1. ..." numbering so items split the same way
        box.paragraphs[0].runs[0].font.size = Pt(20)
    prs.save(path)


def main() -> None:
    PUBLIC.mkdir(parents=True, exist_ok=True)
    for txt in sorted(HERE.glob("*.txt")):
        lines = txt.read_text(encoding="utf-8").splitlines()
        d = docx.Document()
        for line in lines:
            d.add_paragraph(line)
        d.save(txt.with_suffix(".docx"))
        write_pdf(lines, txt.with_suffix(".pdf"))
    write_scanned(HERE / "03_engineering_college_hostel.pdf", HERE / "06_engineering_college_hostel_scanned.pdf")
    write_pptx(HERE / "02_district_hospital_electrification.txt", HERE / "07_district_hospital_presentation.pptx")
    for old in PUBLIC.glob("*"):
        old.unlink()
    for f in HERE.glob("0*.*"):
        shutil.copy(f, PUBLIC / f.name)
    print("Generated:", ", ".join(sorted(p.name for p in HERE.glob("0*.*"))))


if __name__ == "__main__":
    main()
